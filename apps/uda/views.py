from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework import permissions
from django.db import transaction
from django.http import HttpResponse
from django.template.loader import render_to_string
from xhtml2pdf import pisa
import urllib.parse

from .models import Course, UDATemplate, UDA, UDAContent, UDATemplateContent
from .serializers import (
   CourseSerializer, UDATemplateSerializer, UDASerializer,
   UDAContentSerializer, UDATemplateContentSerializer
)
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from io import BytesIO
from django.db.models import Sum
from django.utils.text import get_valid_filename
from bs4 import BeautifulSoup
import logging
import re

logger = logging.getLogger(__name__)

class IsOwnerOrReadOnly(permissions.BasePermission):
    """
    Custom permission to only allow owners of an object to edit it,
    while allowing read-only access for safe methods if authenticated.
    """
    def has_permission(self, request, view):
        return request.user and request.user.is_authenticated

    def has_object_permission(self, request, view, obj):
        if request.method in permissions.SAFE_METHODS:
            return True
        return obj.teacher == request.user


class CourseViewSet(viewsets.ModelViewSet):
    """
    ViewSet per gestire i Corsi (Course).
    Permette CRUD sui corsi e azioni custom per listare e riordinare le UDA.
    """
    queryset = Course.objects.all()
    serializer_class = CourseSerializer
    permission_classes = [IsAuthenticated, IsOwnerOrReadOnly]

    def get_queryset(self):
        user = self.request.user
        if user.is_authenticated:
            return Course.objects.filter(teacher=user)
        return Course.objects.none()

    def perform_create(self, serializer):
        serializer.save(teacher=self.request.user)

    @action(detail=True, methods=['get'], url_path='udas')
    def udas_in_course(self, request, pk=None):
        course = self.get_object()
        udas = UDA.objects.filter(course=course).order_by('order_in_course', 'title')
        serializer = UDASerializer(udas, many=True, context={'request': request})
        return Response(serializer.data)

    @action(detail=True, methods=['post'], url_path='udas/reorder')
    def reorder_udas(self, request, pk=None):
        course = self.get_object()
        uda_ids = request.data.get('uda_ids', [])
        if not isinstance(uda_ids, list):
            return Response({'error': 'uda_ids must be a list.'}, status=status.HTTP_400_BAD_REQUEST)
        current_udas_in_course = UDA.objects.filter(course=course)
        current_uda_ids_map = {uda.id: uda for uda in current_udas_in_course}
        for index, uda_id in enumerate(uda_ids):
            if uda_id not in current_uda_ids_map:
                return Response({'error': f'UDA with id {uda_id} not found in this course or is invalid.'},
                                status=status.HTTP_400_BAD_REQUEST)
            uda_to_update = current_uda_ids_map[uda_id]
            if uda_to_update.order_in_course != index + 1:
                uda_to_update.order_in_course = index + 1
                uda_to_update.save(update_fields=['order_in_course'])
        return Response({'status': 'UDA reordered successfully.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get'], url_path='export-udas-docx')
    def export_udas_docx(self, request, pk=None):
        logger.info(f"EXPORT_UDAS_DOCX CALLED FOR COURSE PK: {pk}")
        course = self.get_object()
        udas_queryset = UDA.objects.filter(course=course).order_by('order_in_course', 'title')
        processed_udas = []
        for uda_instance in udas_queryset:
            total_estimated_hours_val = 0
            for content in uda_instance.contents.all().select_related('lesson'):
                hours_to_add = None
                if content.estimated_hours is not None:
                    hours_to_add = content.estimated_hours
                elif content.content_type == 'LESSON' and content.lesson and content.lesson.estimated_hours is not None:
                    hours_to_add = content.lesson.estimated_hours
                if hours_to_add is not None:
                    try:
                        total_estimated_hours_val += float(hours_to_add)
                    except (ValueError, TypeError):
                        logger.warning(f"Impossibile convertire in float le ore stimate '{hours_to_add}' per il contenuto ID {content.id} dell'UDA ID {uda_instance.id}")
            try:
                if not uda_instance.contents.exists():
                    total_estimated_hours_display = "N/D"
                elif total_estimated_hours_val > 0:
                    total_estimated_hours_display = f"{float(total_estimated_hours_val):.1f}"
                else:
                    total_estimated_hours_display = "0.0"
            except (ValueError, TypeError):
                total_estimated_hours_display = "N/D"
                logger.error(f"Errore imprevisto nella formattazione di total_estimated_hours_display per UDA ID {uda_instance.id}", exc_info=True)
            other_subjects_display = uda_instance.other_involved_subjects_text
            if not other_subjects_display and uda_instance.subjects.exists():
                other_subjects_display = ", ".join([s.name for s in uda_instance.subjects.all()])
            # Processa Lezioni
            lesson_contents_for_template = []
            lesson_contents_queryset = uda_instance.contents.filter(content_type='LESSON').select_related('lesson')
            for lc in lesson_contents_queryset:
                lc_estimated_hours_text = f"{float(lc.estimated_hours):.1f}" if lc.estimated_hours is not None else "N/D"
                lc_actual_hours_text = f"{float(lc.actual_hours):.1f}" if lc.actual_hours is not None else "n/d"
                lesson_contents_for_template.append({
                    'lesson': lc.lesson,
                    'estimated_hours_display': lc_estimated_hours_text,
                    'actual_hours_display': lc_actual_hours_text,
                })

            # Processa Attività
            activity_contents_for_template = []
            activity_contents_queryset = uda_instance.contents.filter(content_type='ACTIVITY')
            for ac in activity_contents_queryset:
                ac_estimated_hours_text = f"{float(ac.estimated_hours):.1f}" if ac.estimated_hours is not None else "N/D"
                ac_actual_hours_text = f"{float(ac.actual_hours):.1f}" if ac.actual_hours is not None else "n/d"
                activity_contents_for_template.append({
                    'title': ac.activity_title or "Attività senza titolo",
                    'estimated_hours_display': ac_estimated_hours_text,
                    'actual_hours_display': ac_actual_hours_text,
                })
            
            # Processa Note
            note_contents_for_template = []
            note_contents_queryset = uda_instance.contents.filter(content_type='NOTE')
            for nc in note_contents_queryset:
                note_contents_for_template.append({
                    'title': nc.note_title or "Nota senza titolo",
                    'description': nc.note_content or "",
                })

            processed_udas.append({
                'title': uda_instance.title,
                'start_date': uda_instance.start_date.strftime('%d/%m/%Y') if uda_instance.start_date else 'N/D',
                'end_date': uda_instance.end_date.strftime('%d/%m/%Y') if uda_instance.end_date else 'N/D',
                'description': uda_instance.description,
                'competences_html': uda_instance.competences_html,
                'topics_string': ", ".join([topic.name for topic in uda_instance.topics.all()]),
                'knowledge_html': uda_instance.knowledge_html,
                'skills_html': uda_instance.skills_html,
                'total_estimated_hours_display': total_estimated_hours_display,
                'didactic_strategies_html': uda_instance.didactic_strategies_html,
                'materials_tools_html': uda_instance.materials_tools_html,
                'assessment_type_html': uda_instance.assessment_type_html,
                'evaluation_html': uda_instance.evaluation_html,
                'other_involved_subjects_display': other_subjects_display or "N/D",
                'is_civic_education': uda_instance.is_civic_education,
                'lesson_contents': lesson_contents_for_template,
                'activity_contents': activity_contents_for_template,
                'note_contents': note_contents_for_template,
                'export_specific_annotations_html': uda_instance.export_specific_annotations_html,
            })

        # Sanitizzazione del nome del corso più robusta
        course_name_base = course.name
        course_name_sanitized = get_valid_filename(course_name_base)
        
        if not course_name_sanitized:
            # Se get_valid_filename fallisce sul nome originale, prova con una versione solo alfanumerica
            course_name_alphanumeric = re.sub(r'[^a-zA-Z0-9]', '', course_name_base)
            if course_name_alphanumeric: # Solo se la versione alfanumerica non è vuota
                course_name_sanitized = get_valid_filename(course_name_alphanumeric)
        
        if not course_name_sanitized: # Fallback finale se ancora vuoto
            course_name_sanitized = "corso"
        
        # Rimuove eventuali underscore finali dal risultato
        while course_name_sanitized.endswith("_"):
            course_name_sanitized = course_name_sanitized[:-1]
        
        if not course_name_sanitized: # Fallback se dopo la rimozione degli underscore è vuoto
            course_name_sanitized = "corso"

        filename_docx_raw = f"export_udas_{course_name_sanitized}.docx"
        filename_docx_encoded = urllib.parse.quote(filename_docx_raw)

        document = Document()
        document.add_heading(f'Programmazione Didattica Individuale - UDA del Corso: {course.name}', level=1)
        p_docente = document.add_paragraph()
        p_docente.add_run(f"Docente: {request.user.get_full_name() or request.user.username}").bold = True
        def add_table_row_docx(table, label, value_html_or_text, is_html=True):
            row_cells = table.add_row().cells
            row_cells[0].text = label
            processed_value = "N/D"
            if value_html_or_text:
                if is_html:
                    soup = BeautifulSoup(str(value_html_or_text), 'html.parser')
                    processed_value = soup.get_text(separator='\n', strip=True) or "N/D"
                else:
                    processed_value = str(value_html_or_text)
            row_cells[1].text = processed_value
            row_cells[0].paragraphs[0].runs[0].bold = True
        for index, uda_data in enumerate(processed_udas):
            document.add_heading(f'UDA n. {index + 1}: {uda_data["title"]}', level=2)
            
            p_periodo = document.add_paragraph()
            p_periodo.add_run(f"Periodo: dal {uda_data['start_date']} al {uda_data['end_date']}").italic = True
            
            table_details = document.add_table(rows=0, cols=2)
            table_details.style = 'TableGrid'
            table_details.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_details.columns[0].width = Inches(2.0)
            table_details.columns[1].width = Inches(4.5)

            add_table_row_docx(table_details, 'Descrizione', uda_data['description'], is_html=True)
            add_table_row_docx(table_details, 'Tempi (durata in ore)', f"{uda_data['total_estimated_hours_display']} ore", is_html=False)
            add_table_row_docx(table_details, 'Argomenti Uda', uda_data['topics_string'], is_html=False)
            add_table_row_docx(table_details, 'Competenze attese a livello di UdA', uda_data['competences_html'])
            add_table_row_docx(table_details, 'Conoscenze ADA (sapere)', uda_data['knowledge_html'])
            add_table_row_docx(table_details, 'Abilità-Capacità ADA (saper fare)', uda_data['skills_html'])
            add_table_row_docx(table_details, 'Strategie didattiche', uda_data['didactic_strategies_html'])
            add_table_row_docx(table_details, 'Materiali e strumenti', uda_data['materials_tools_html'])
            add_table_row_docx(table_details, 'Tipo di verifiche', uda_data['assessment_type_html'])
            add_table_row_docx(table_details, 'Valutazione', uda_data['evaluation_html'])
            add_table_row_docx(table_details, 'Discipline Coinvolte', uda_data['other_involved_subjects_display'], is_html=False)
            
            if uda_data['export_specific_annotations_html']:
                add_table_row_docx(table_details, 'Annotazioni', uda_data['export_specific_annotations_html'], is_html=True)

            if uda_data['is_civic_education']:
                add_table_row_docx(table_details, 'Educazione Civica', 'Sì, parte del percorso', is_html=False)

            if uda_data['lesson_contents']:
                document.add_heading('Lezioni Previste', level=3)
                table_lessons = document.add_table(rows=1, cols=3)
                table_lessons.style = 'TableGrid'
                table_lessons.alignment = WD_TABLE_ALIGNMENT.CENTER
                hdr_cells = table_lessons.rows[0].cells
                hdr_cells[0].text = 'Titolo Lezione'
                hdr_cells[1].text = 'Durata Stimata (ore)'
                hdr_cells[2].text = 'Durata Effettiva (ore)'
                for cell in hdr_cells:
                    cell.paragraphs[0].runs[0].bold = True
                for lc_data in uda_data['lesson_contents']:
                    row_cells_lc = table_lessons.add_row().cells
                    row_cells_lc[0].text = lc_data['lesson'].title if lc_data['lesson'] else "Lezione non specificata"
                    row_cells_lc[1].text = lc_data['estimated_hours_display']
                    row_cells_lc[2].text = lc_data['actual_hours_display']

            if uda_data['activity_contents']:
                document.add_heading('Attività Previste', level=3)
                table_activities = document.add_table(rows=1, cols=3)
                table_activities.style = 'TableGrid'
                table_activities.alignment = WD_TABLE_ALIGNMENT.CENTER
                hdr_cells_ac = table_activities.rows[0].cells
                hdr_cells_ac[0].text = 'Titolo Attività'
                hdr_cells_ac[1].text = 'Ore Stimate'
                hdr_cells_ac[2].text = 'Ore Effettive'
                for cell in hdr_cells_ac:
                    cell.paragraphs[0].runs[0].bold = True
                
                for ac_data in uda_data['activity_contents']:
                    row_cells_ac = table_activities.add_row().cells
                    row_cells_ac[0].text = ac_data['title']
                    row_cells_ac[1].text = ac_data['estimated_hours_display']
                    row_cells_ac[2].text = ac_data['actual_hours_display']

            if uda_data['note_contents']:
                document.add_heading('Note', level=3)
                for note_data in uda_data['note_contents']:
                    document.add_paragraph(note_data['title'], style='Intense Quote') # Usa uno stile per il titolo della nota
                    # Pulisci e aggiungi la descrizione HTML
                    if note_data['description']:
                        soup = BeautifulSoup(note_data['description'], 'html.parser')
                        # Aggiungi il testo pulito, preservando i paragrafi
                        for p in soup.find_all(['p', 'li']):
                             document.add_paragraph(p.get_text(strip=True), style='List Paragraph')
                    document.add_paragraph() # Spazio dopo ogni nota

            document.add_paragraph()
            if index < len(processed_udas) - 1:
                document.add_page_break()
        file_stream = BytesIO()
        document.save(file_stream)
        file_stream.seek(0)
        response = HttpResponse(
            file_stream.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename_docx_encoded}"
        return response

    @action(detail=True, methods=['get'], url_path='export-udas-pdf')
    def export_udas_pdf(self, request, pk=None):
        logger.info(f"EXPORT_UDAS_PDF CALLED FOR COURSE PK: {pk}")
        course = self.get_object()
        udas_queryset = UDA.objects.filter(course=course).order_by('order_in_course', 'title')
        processed_udas = []
        for uda_instance in udas_queryset:
            total_estimated_hours_val = uda_instance.contents.aggregate(Sum('estimated_hours'))['estimated_hours__sum'] or 0
            try:
                total_estimated_hours_display = f"{float(total_estimated_hours_val):.1f}"
            except (ValueError, TypeError):
                total_estimated_hours_display = "N/D"
            other_subjects_display = uda_instance.other_involved_subjects_text
            if not other_subjects_display and uda_instance.subjects.exists():
                other_subjects_display = ", ".join([s.name for s in uda_instance.subjects.all()])
            annotations_text_for_pdf = uda_instance.description
            is_html_annotations = False
            if uda_instance.export_specific_annotations_html:
                annotations_text_for_pdf = uda_instance.export_specific_annotations_html
                is_html_annotations = True
            # Processa Lezioni
            lesson_contents_for_template = []
            lesson_contents_queryset = uda_instance.contents.filter(content_type='LESSON').select_related('lesson')
            for lc in lesson_contents_queryset:
                lc_hours_text = f"{float(lc.estimated_hours):.1f}" if lc.estimated_hours is not None else "N/D"
                lesson_contents_for_template.append({
                    'lesson': lc.lesson,
                    'hours_display_pdf': lc_hours_text
                })

            # Processa Attività
            activity_contents_for_template = []
            activity_contents_queryset = uda_instance.contents.filter(content_type='ACTIVITY')
            for ac in activity_contents_queryset:
                ac_estimated_hours_text = f"{float(ac.estimated_hours):.1f}" if ac.estimated_hours is not None else "N/D"
                ac_actual_hours_text = f"{float(ac.actual_hours):.1f}" if ac.actual_hours is not None else "N/D"
                activity_contents_for_template.append({
                    'title': ac.activity_title or "Attività senza titolo",
                    'estimated_hours_display': ac_estimated_hours_text,
                    'actual_hours_display': ac_actual_hours_text,
                })
            processed_udas.append({
                'title': uda_instance.title,
                'competences_html': uda_instance.competences_html,
                'topics_string': ", ".join([topic.name for topic in uda_instance.topics.all()]),
                'knowledge_html': uda_instance.knowledge_html,
                'skills_html': uda_instance.skills_html,
                'total_estimated_hours_display': total_estimated_hours_display,
                'didactic_strategies_html': uda_instance.didactic_strategies_html,
                'materials_tools_html': uda_instance.materials_tools_html,
                'assessment_type_html': uda_instance.assessment_type_html,
                'evaluation_html': uda_instance.evaluation_html,
                'other_involved_subjects_display': other_subjects_display or "N/D",
                'annotations_display': annotations_text_for_pdf or "N/D",
                'is_html_annotations': is_html_annotations,
                'is_civic_education': uda_instance.is_civic_education,
                'lesson_contents': lesson_contents_for_template,
                'activity_contents': activity_contents_for_template,
            })

        # Sanitizzazione del nome del corso più robusta
        course_name_base = course.name
        course_name_sanitized = get_valid_filename(course_name_base)

        if not course_name_sanitized:
             # Se get_valid_filename fallisce sul nome originale, prova con una versione solo alfanumerica
            course_name_alphanumeric = re.sub(r'[^a-zA-Z0-9]', '', course_name_base)
            if course_name_alphanumeric: # Solo se la versione alfanumerica non è vuota
                course_name_sanitized = get_valid_filename(course_name_alphanumeric)

        if not course_name_sanitized: # Fallback finale se ancora vuoto
            course_name_sanitized = "corso"
            
        # Rimuove eventuali underscore finali dal risultato
        while course_name_sanitized.endswith("_"):
            course_name_sanitized = course_name_sanitized[:-1]
            
        if not course_name_sanitized: # Fallback se dopo la rimozione degli underscore è vuoto
            course_name_sanitized = "corso"

        filename_pdf_raw = f"export_udas_{course_name_sanitized}.pdf"
        filename_pdf_encoded = urllib.parse.quote(filename_pdf_raw)

        try:
            context = {
                'course': course,
                'udas': processed_udas,
                'teacher_name': request.user.get_full_name() or request.user.username
            }
            html_string = render_to_string('uda_export/uda_pdf_template.html', context)
            pdf_file_stream = BytesIO()
            pdf = pisa.CreatePDF(
                html_string,
                dest=pdf_file_stream,
                encoding='UTF-8'
            )
            pdf_file_stream.seek(0)
            if pdf.err:
                error_messages = "; ".join([str(msg) for msg in pdf.log if msg.level >= pisa.ERROR]) if hasattr(pdf, 'log') else str(pdf.err)
                logger.error(f"Errore pisa.CreatePDF per corso {course.id}: codice {pdf.err} - {error_messages}")
                return HttpResponse(f"Errore durante la generazione del PDF (codice: {pdf.err})", status=500)
            response = HttpResponse(pdf_file_stream.read(), content_type='application/pdf')
            response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename_pdf_encoded}"
            return response
        except Exception as e:
            logger.error(f"Eccezione durante la generazione del PDF per corso {course.id}: {e}", exc_info=True)
            return HttpResponse(f"Errore interno del server durante la generazione del PDF: {str(e)}", status=500)

    @action(detail=True, methods=['post'], url_path='copy')
    @transaction.atomic
    def copy(self, request, pk=None):
        """
        Crea una copia di un Corso, incluse tutte le sue UDA.
        """
        original_course = self.get_object()
        
        # 1. Crea la copia del corso
        new_course = Course.objects.create(
            teacher=request.user,
            name=f"Copia di {original_course.name}",
            description=original_course.description
        )
        
        # 2. Recupera e copia le UDA
        udas_to_copy = original_course.udas.all().order_by('order_in_course')
        for original_uda in udas_to_copy:
            # 3. Usa la funzione helper per copiare ogni UDA e associarla al nuovo corso
            _copy_uda_instance(original_uda, new_course, request.user)
            
        serializer = self.get_serializer(new_course)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

class UDATemplateViewSet(viewsets.ModelViewSet):
    """
    ViewSet per gestire i Template UDA (UDATemplate).
    Permette CRUD sui template e gestione dei loro contenuti.
    """
    queryset = UDATemplate.objects.all()
    serializer_class = UDATemplateSerializer
    permission_classes = [IsAuthenticated, IsOwnerOrReadOnly]

    def get_queryset(self):
        user = self.request.user
        if user.is_authenticated:
            return UDATemplate.objects.filter(teacher=user)
        return UDATemplate.objects.none()

    def perform_create(self, serializer):
        serializer.save()

    @action(detail=True, methods=['get', 'post'], serializer_class=UDATemplateContentSerializer, url_path='contents')
    def contents(self, request, pk=None):
        template = self.get_object()
        if request.method == 'GET':
            contents = template.contents.all()
            serializer = UDATemplateContentSerializer(contents, many=True)
            return Response(serializer.data)
        elif request.method == 'POST':
            serializer = UDATemplateContentSerializer(data=request.data)
            if serializer.is_valid():
                serializer.save(uda_template=template)
                return Response(serializer.data, status=status.HTTP_201_CREATED)
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['get', 'put', 'patch', 'delete'],
            serializer_class=UDATemplateContentSerializer,
            url_path='contents/(?P<content_pk>[^/.]+)')
    def single_content(self, request, pk=None, content_pk=None):
        template = self.get_object()
        try:
            content = template.contents.get(pk=content_pk)
        except UDATemplateContent.DoesNotExist:
            return Response({'error': 'Content not found.'}, status=status.HTTP_404_NOT_FOUND)
        if request.method == 'GET':
            serializer = UDATemplateContentSerializer(content)
            return Response(serializer.data)
        elif request.method in ['PUT', 'PATCH']:
            serializer = UDATemplateContentSerializer(content, data=request.data, partial=request.method == 'PATCH')
            if serializer.is_valid():
                serializer.save()
                return Response(serializer.data)
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        elif request.method == 'DELETE':
            content.delete()
            return Response(status=status.HTTP_204_NO_CONTENT)


def _copy_uda_instance(original_uda, new_course, teacher):
    """
    Funzione helper per copiare un'istanza di UDA.
    Crea una copia profonda dell'UDA, dei suoi contenuti e delle relazioni M2M.
    Associa la nuova UDA al corso specificato.
    """
    new_uda = UDA.objects.create(
        teacher=teacher,
        title=f"Copia di {original_uda.title}",
        description=original_uda.description,
        start_date=original_uda.start_date,
        end_date=original_uda.end_date,
        status='TODO',
        course=new_course,  # Associa al nuovo corso passato come argomento
        order_in_course=original_uda.order_in_course if new_course else None,
        source_template=None,
        is_civic_education=original_uda.is_civic_education,
        didactic_strategies_html=original_uda.didactic_strategies_html,
        materials_tools_html=original_uda.materials_tools_html,
        assessment_type_html=original_uda.assessment_type_html,
        evaluation_html=original_uda.evaluation_html,
        other_involved_subjects_text=original_uda.other_involved_subjects_text,
        export_specific_annotations_html=original_uda.export_specific_annotations_html,
        competences_html=original_uda.competences_html,
        knowledge_html=original_uda.knowledge_html,
        skills_html=original_uda.skills_html
    )
    new_uda.subjects.set(original_uda.subjects.all())
    new_uda.topics.set(original_uda.topics.all())

    original_contents = original_uda.contents.all().order_by('order')
    for original_content in original_contents:
        UDAContent.objects.create(
            uda=new_uda,
            content_type=original_content.content_type,
            lesson=original_content.lesson,
            quiz_template=original_content.quiz_template,
            note_title=original_content.note_title,
            note_content=original_content.note_content,
            activity_title=original_content.activity_title,
            activity_description=original_content.activity_description,
            activity_attachment_url=original_content.activity_attachment_url,
            activity_completed=False,
            teacher_marked_completed=False,
            order=original_content.order,
            estimated_hours=original_content.estimated_hours
        )
    return new_uda


class UDAViewSet(viewsets.ModelViewSet):
    """
    ViewSet per gestire le UDA.
    Permette CRUD sulle UDA e gestione dei loro contenuti, incluso il completamento delle attività.
    """
    queryset = UDA.objects.all()
    serializer_class = UDASerializer
    permission_classes = [IsAuthenticated, IsOwnerOrReadOnly]
 
    def get_queryset(self):
        user = self.request.user
        if user.is_authenticated:
            status_filter = self.request.query_params.get('status')
            course_filter = self.request.query_params.get('course_id')
            queryset = UDA.objects.filter(teacher=user)
            if status_filter:
                queryset = queryset.filter(status=status_filter)
            if course_filter:
               try:
                   course_id = int(course_filter)
                   queryset = queryset.filter(course_id=course_id)
               except ValueError:
                   pass
            queryset = queryset.order_by('course__name', 'order_in_course', 'title')
            return queryset
        return UDA.objects.none()

    def perform_create(self, serializer):
        serializer.save()

    @action(detail=True, methods=['get', 'post'], serializer_class=UDAContentSerializer, url_path='contents')
    def contents(self, request, pk=None):
        uda = self.get_object()
        if request.method == 'GET':
            contents = uda.contents.all()
            serializer = UDAContentSerializer(contents, many=True, context={'request': request})
            return Response(serializer.data)
        elif request.method == 'POST':
            serializer = UDAContentSerializer(data=request.data)
            if serializer.is_valid():
                serializer.save(uda=uda)
                return Response(serializer.data, status=status.HTTP_201_CREATED)
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['get', 'put', 'patch', 'delete'],
            serializer_class=UDAContentSerializer,
            url_path='contents/(?P<content_pk>[^/.]+)')
    def single_content(self, request, pk=None, content_pk=None):
        uda = self.get_object()
        try:
            content = uda.contents.get(pk=content_pk)
        except UDAContent.DoesNotExist:
            return Response({'error': 'Content not found.'}, status=status.HTTP_404_NOT_FOUND)

        if request.method == 'GET':
            serializer = UDAContentSerializer(content, context={'request': request})
            return Response(serializer.data)
        elif request.method in ['PUT', 'PATCH']:
            serializer = UDAContentSerializer(content, data=request.data, partial=request.method == 'PATCH', context={'request': request})
            if serializer.is_valid():
                serializer.save()
                return Response(serializer.data)
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        elif request.method == 'DELETE':
            content.delete()
            return Response(status=status.HTTP_204_NO_CONTENT)

    @action(detail=True, methods=['patch'],
            url_path='contents/(?P<content_pk>[^/.]+)/complete-activity')
    def complete_activity(self, request, pk=None, content_pk=None):
        uda = self.get_object()
        try:
            content = uda.contents.get(pk=content_pk, content_type='ACTIVITY')
        except UDAContent.DoesNotExist:
            return Response({'error': 'Activity content not found.'}, status=status.HTTP_404_NOT_FOUND)
        
        content.activity_completed = request.data.get('completed', False)
        content.save(update_fields=['activity_completed'])
        return Response({'status': 'Activity completion status updated.'})

    @action(detail=True, methods=['patch'],
            url_path='contents/(?P<content_pk>[^/.]+)/update-teacher-completion')
    def update_teacher_completion(self, request, pk=None, content_pk=None):
        uda = self.get_object()
        try:
            content = uda.contents.get(pk=content_pk)
        except UDAContent.DoesNotExist:
            return Response({'error': 'Content not found.'}, status=status.HTTP_404_NOT_FOUND)
        
        completed = request.data.get('completed')
        if completed is None:
            return Response({'error': 'Missing "completed" field.'}, status=status.HTTP_400_BAD_REQUEST)
        
        content.teacher_marked_completed = bool(completed)
        content.save(update_fields=['teacher_marked_completed'])
        return Response({'status': 'Teacher completion status updated.'})

    @action(detail=True, methods=['post'], url_path='contents/reorder')
    def reorder_contents(self, request, pk=None):
        """
        Riordina i contenuti di una UDA.
        Si aspetta una lista di ID di contenuti nell'ordine desiderato.
        Esempio body: { "content_ids": [3, 1, 2] }
        """
        uda = self.get_object()
        content_ids = request.data.get('content_ids', [])
        
        if not isinstance(content_ids, list):
            return Response({'error': 'content_ids must be a list.'}, status=status.HTTP_400_BAD_REQUEST)
        
        with transaction.atomic():
            current_contents = {content.id: content for content in uda.contents.all()}
            
            if len(content_ids) != len(current_contents):
                return Response({'error': 'The list of IDs does not match the number of contents in the UDA.'}, status=status.HTTP_400_BAD_REQUEST)

            for index, content_id in enumerate(content_ids):
                if content_id not in current_contents:
                    return Response({'error': f'Content with id {content_id} not found in this UDA.'}, status=status.HTTP_400_BAD_REQUEST)
                
                content_to_update = current_contents[content_id]
                if content_to_update.order != index:
                    content_to_update.order = index
                    content_to_update.save(update_fields=['order'])

        return Response({'status': 'Contents reordered successfully.'}, status=status.HTTP_200_OK)

    @action(detail=True, methods=['post'], url_path='copy')
    @transaction.atomic
    def copy_uda(self, request, pk=None):
        """
        Crea una copia di una UDA esistente, inclusi i suoi contenuti.
        La nuova UDA non sarà associata a nessun corso.
        """
        original_uda = self.get_object()
        # Chiama la funzione helper per creare la copia
        new_uda = _copy_uda_instance(original_uda, new_course=None, teacher=request.user)
        
        serializer = self.get_serializer(new_uda)
        return Response(serializer.data, status=status.HTTP_201_CREATED)
